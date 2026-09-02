from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import colors as mpl_colors
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "supplementary"
    / "machine_readable"
    / "table_s13_chemical_subgroup_delta_rmse_all_96_rows.csv"
)
OUTPUT_DIR = ROOT / "figures" / "generated"
PPTX = OUTPUT_DIR / "figure_2_chemical_subgroup_gain_editable.pptx"
PREVIEW = OUTPUT_DIR / "figure_2_chemical_subgroup_gain.png"
PDF = OUTPUT_DIR / "figure_2_chemical_subgroup_gain.pdf"

PANEL_IDS = ["R01", "R02", "R03", "R04", "R05", "R09"]
PANEL_LABELS = [
    "AqSolDBc\nOOF", "ComPlat\nOOF", "TDC\nOOF",
    "TDC\nsupplied\ntest", "JCheM\nsupplied\ntest", "ComPlat\nsupplied\ntest",
]
AXIS_GROUPS = [
    ("Molecular\nweight", ["<250 Da", "250–<400 Da", "≥400 Da"]),
    ("Rotatable\nbonds", ["0", "1–4", "≥5"]),
    ("H-bond\ndonor", ["No", "Yes"]),
    ("H-bond\nacceptor", ["No", "Yes"]),
    ("Aromatic\nring", ["No", "Yes"]),
    ("N-containing\nheterocycle", ["No", "Yes"]),
    ("Carboxyl\ngroup", ["No", "Yes"]),
]

SLIDE_W = 17.4
SLIDE_H = 9.65
GRID_Y = 1.40
CELL_H = 0.43
CELL_W = 0.84
GRID_H = CELL_H * 16
GRID_W = CELL_W * 6
PANEL_A_X = 3.20
PANEL_B_X = 8.72


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value.replace("#", "").upper())


def to_hex(rgba) -> str:
    r, g, b = [round(255 * x) for x in rgba[:3]]
    return f"{r:02X}{g:02X}{b:02X}"


def add_text(slide, name, text, x, y, w, h, size=8, bold=False,
             color="1F2937", align=PP_ALIGN.CENTER, rotation=0):
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    box.name = name
    box.rotation = rotation
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Cm(0.02)
    tf.margin_right = Cm(0.02)
    tf.margin_top = Cm(0.01)
    tf.margin_bottom = Cm(0.01)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = 0.9
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rect(slide, name, x, y, w, h, fill, line="FFFFFF", line_width=0.45):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h)
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(line_width)
    return shape


def add_line(slide, name, x1, y1, x2, y2, color="374151", width=0.8):
    line = slide.shapes.add_connector(1, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    line.name = name
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def load_matrices():
    effects = pd.read_csv(SOURCE)
    matrices, ns = {}, {}
    for comparator in ("xgb", "neural"):
        matrix = np.full((16, 6), np.nan)
        n_matrix = np.zeros((16, 6), dtype=int)
        for i in range(16):
            for j, panel_id in enumerate(PANEL_IDS):
                row = effects[
                    (effects.level_index == i) &
                    (effects.evaluation_id == panel_id)
                ].iloc[0]
                matrix[i, j] = row[f"delta_rmse_{comparator}"]
                n_matrix[i, j] = int(row.n)
        matrices[comparator] = matrix
        ns[comparator] = n_matrix
    vmax = float(np.nanmax(np.abs(np.concatenate([
        matrices["xgb"].ravel(), matrices["neural"].ravel()
    ]))))
    return matrices, ns, vmax


def build_ppt():
    matrices, ns, vmax = load_matrices()
    cmap = matplotlib.colormaps["RdBu_r"]
    norm = mpl_colors.Normalize(vmin=-vmax, vmax=vmax)

    prs = Presentation()
    prs.slide_width = Cm(SLIDE_W)
    prs.slide_height = Cm(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")

    add_text(slide, "figure_title",
             "DLG-Sol gain across prespecified chemical subgroups",
             0.25, 0.18, 16.90, 0.48, 12, True)
    add_text(slide, "delta_definition",
             "ΔRMSE = RMSE(DLG-Sol) − RMSE(component); negative values favor DLG-Sol",
             0.25, 0.67, 16.90, 0.32, 7.4, False, "4B5563")

    add_text(slide, "row_header_axis", "Chemical axis",
             0.10, 1.06, 1.45, 0.28, 7.0, True, "4B5563", PP_ALIGN.LEFT)
    add_text(slide, "row_header_level", "Subgroup",
             1.76, 1.06, 1.30, 0.28, 7.0, True, "4B5563", PP_ALIGN.RIGHT)

    row_cursor = 0
    group_boundaries = []
    for gidx, (axis_name, levels) in enumerate(AXIS_GROUPS):
        n_levels = len(levels)
        group_y = GRID_Y + row_cursor * CELL_H
        add_rect(slide, f"axis_band_{gidx+1}", 0.14, group_y,
                 3.00, n_levels * CELL_H,
                 "F8FAFC" if gidx % 2 == 0 else "FFFFFF", "FFFFFF", 0)
        add_text(slide, f"axis_name_{gidx+1}", axis_name,
                 0.10, group_y, 1.45, n_levels * CELL_H,
                 6.8, True, "1F2937", PP_ALIGN.LEFT)
        for lidx, level in enumerate(levels):
            y = group_y + lidx * CELL_H
            add_text(slide, f"axis_{gidx+1}_level_{lidx+1}", level,
                     1.76, y, 1.30, CELL_H, 6.6, False,
                     "374151", PP_ALIGN.RIGHT)
        row_cursor += n_levels
        if row_cursor < 16:
            group_boundaries.append(row_cursor)

    panel_specs = [
        ("A", "xgb", "DLG-Sol − XGB", PANEL_A_X),
        ("B", "neural", "DLG-Sol − neural component", PANEL_B_X),
    ]
    for panel_letter, comparator, title, panel_x in panel_specs:
        add_text(slide, f"panel_{panel_letter}_title",
                 f"{panel_letter}   {title}", panel_x, 1.00,
                 GRID_W, 0.32, 9.5, True, "111827", PP_ALIGN.LEFT)
        matrix = matrices[comparator]
        n_matrix = ns[comparator]
        for i in range(16):
            for j in range(6):
                value = float(matrix[i, j])
                fill = to_hex(cmap(norm(value)))
                x = panel_x + j * CELL_W
                y = GRID_Y + i * CELL_H
                name = (
                    f"panel_{panel_letter}_{PANEL_IDS[j]}_row_{i+1}_"
                    f"delta_{value:+.5f}_n_{n_matrix[i,j]}"
                )
                cell = add_rect(slide, name, x, y, CELL_W, CELL_H,
                                fill, "FFFFFF", 0.5)
                cell.text_frame.clear()

        border = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Cm(panel_x), Cm(GRID_Y),
            Cm(GRID_W), Cm(GRID_H)
        )
        border.name = f"panel_{panel_letter}_outer_border"
        border.fill.background()
        border.line.color.rgb = rgb("374151")
        border.line.width = Pt(0.9)
        for boundary in group_boundaries:
            y = GRID_Y + boundary * CELL_H
            add_line(slide, f"panel_{panel_letter}_group_separator_{boundary}",
                     panel_x, y, panel_x + GRID_W, y, "374151", 0.85)

        for j, label in enumerate(PANEL_LABELS):
            add_text(slide, f"panel_{panel_letter}_column_{j+1}", label,
                     panel_x + j * CELL_W - 0.30, GRID_Y + GRID_H + 0.02,
                     1.26, 0.56, 6.0, False, "1F2937",
                     PP_ALIGN.RIGHT, rotation=315)

    cbar_x = 14.25
    cbar_y = 2.05
    cbar_w = 0.34
    cbar_h = 4.95
    steps = 30
    for k in range(steps):
        frac = k / (steps - 1)
        value = vmax - frac * 2 * vmax
        add_rect(slide, f"colorbar_step_{k+1}", cbar_x,
                 cbar_y + k * cbar_h / steps,
                 cbar_w, cbar_h / steps + 0.005,
                 to_hex(cmap(norm(value))), to_hex(cmap(norm(value))), 0)
    border = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Cm(cbar_x), Cm(cbar_y), Cm(cbar_w), Cm(cbar_h)
    )
    border.name = "colorbar_border"
    border.fill.background()
    border.line.color.rgb = rgb("374151")
    border.line.width = Pt(0.8)
    ticks = [vmax, 0.10, 0.05, 0.00, -0.05, -0.10, -vmax]
    for idx, value in enumerate(ticks):
        frac = (vmax - value) / (2 * vmax)
        y = cbar_y + frac * cbar_h
        add_line(slide, f"colorbar_tick_{idx+1}", cbar_x + cbar_w,
                 y, cbar_x + cbar_w + 0.10, y, "374151", 0.6)
        label = f"{value:+.2f}" if abs(value) > 1e-12 else "0"
        add_text(slide, f"colorbar_tick_label_{idx+1}", label,
                 cbar_x + cbar_w + 0.12, y - 0.16, 0.60, 0.32,
                 6.8, False, "374151", PP_ALIGN.LEFT)
    add_text(slide, "colorbar_label", "ΔRMSE",
             14.02, 1.62, 1.10, 0.32, 8.2, True, "111827", PP_ALIGN.LEFT)
    add_text(slide, "colorbar_high_label", "DLG-Sol\nhigher",
             15.55, 2.00, 1.35, 0.58, 6.5, False, "374151", PP_ALIGN.LEFT)
    add_text(slide, "colorbar_low_label", "DLG-Sol\nlower",
             15.55, 6.48, 1.35, 0.58, 6.5, False, "374151", PP_ALIGN.LEFT)

    prs.core_properties.title = "Figure 2 — DLG-Sol gain across chemical subgroups"
    prs.core_properties.subject = "Editable two-panel heatmap with 192 native PowerPoint cells"
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.core_properties.comments = (
        "All cells, labels, borders, separators, and colour-bar steps are editable. "
        "Source: supplementary/machine_readable/"
        "table_s13_chemical_subgroup_delta_rmse_all_96_rows.csv. "
        f"Symmetric colour range: {-vmax:.8f} to {vmax:.8f}."
    )
    prs.save(PPTX)
    return vmax


def build_preview(vmax):
    matrices, _, _ = load_matrices()
    fig = plt.figure(
        figsize=(SLIDE_W / 2.54, SLIDE_H / 2.54), dpi=300, facecolor="white"
    )
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, SLIDE_W)
    ax.set_ylim(SLIDE_H, 0)
    ax.axis("off")
    cmap = matplotlib.colormaps["RdBu_r"]
    norm = mpl_colors.Normalize(vmin=-vmax, vmax=vmax)
    ax.text(SLIDE_W / 2, 0.42,
            "DLG-Sol gain across prespecified chemical subgroups",
            ha="center", va="center", fontsize=12, fontweight="bold", color="#1F2937")
    ax.text(SLIDE_W / 2, 0.82,
            "ΔRMSE = RMSE(DLG-Sol) − RMSE(component); negative values favor DLG-Sol",
            ha="center", va="center", fontsize=7.4, color="#4B5563")

    cursor = 0
    for gidx, (axis_name, levels) in enumerate(AXIS_GROUPS):
        gy = GRID_Y + cursor * CELL_H
        gh = len(levels) * CELL_H
        if gidx % 2 == 0:
            ax.add_patch(plt.Rectangle((0.14, gy), 3.0, gh, color="#F8FAFC", ec="none"))
        ax.text(0.10, gy + gh / 2, axis_name, ha="left", va="center",
                fontsize=6.2, fontweight="bold", color="#1F2937", linespacing=0.9)
        for lidx, level in enumerate(levels):
            ax.text(3.02, gy + (lidx + 0.5) * CELL_H, level,
                    ha="right", va="center", fontsize=6.2, color="#374151")
        cursor += len(levels)

    for letter, comp, title, px in [
        ("A", "xgb", "DLG-Sol − XGB", PANEL_A_X),
        ("B", "neural", "DLG-Sol − neural component", PANEL_B_X),
    ]:
        ax.text(px, 1.15, f"{letter}   {title}", ha="left", va="center",
                fontsize=9.5, fontweight="bold", color="#111827")
        for i in range(16):
            for j in range(6):
                x, y = px + j * CELL_W, GRID_Y + i * CELL_H
                ax.add_patch(plt.Rectangle(
                    (x, y), CELL_W, CELL_H,
                    facecolor=cmap(norm(matrices[comp][i, j])),
                    edgecolor="white", linewidth=0.6
                ))
        ax.add_patch(plt.Rectangle((px, GRID_Y), GRID_W, GRID_H,
                                   fill=False, ec="#374151", lw=0.9))
        for b in (3, 6, 8, 10, 12, 14):
            ax.plot([px, px + GRID_W], [GRID_Y + b * CELL_H] * 2,
                    color="#374151", lw=0.8)
        for j, label in enumerate(PANEL_LABELS):
            ax.text(px + (j + 0.58) * CELL_W, GRID_Y + GRID_H + 0.12,
                    label, ha="right", va="top", rotation=45,
                    rotation_mode="anchor", fontsize=5.8,
                    color="#1F2937", linespacing=0.85)

    sm = plt.cm.ScalarMappable(norm=norm, cmap=cmap)
    cax = fig.add_axes([14.25 / SLIDE_W, (SLIDE_H - 7.0) / SLIDE_H,
                        0.34 / SLIDE_W, 4.95 / SLIDE_H])
    cb = fig.colorbar(sm, cax=cax)
    cb.ax.tick_params(labelsize=6)
    ax.text(14.02, 1.84, "ΔRMSE", ha="left", fontsize=8.2, fontweight="bold")
    ax.text(15.55, 2.05, "DLG-Sol\nhigher", ha="left", va="top",
            fontsize=6.5, color="#374151", linespacing=0.9)
    ax.text(15.55, 6.55, "DLG-Sol\nlower", ha="left", va="top",
            fontsize=6.5, color="#374151", linespacing=0.9)
    fig.savefig(PREVIEW, facecolor="white")
    fig.savefig(PDF, facecolor="white")
    plt.close(fig)


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    vmax = build_ppt()
    build_preview(vmax)
    print(PPTX.relative_to(ROOT))
    print(PREVIEW.relative_to(ROOT))
    print(PDF.relative_to(ROOT))
